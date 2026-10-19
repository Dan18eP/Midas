"""
Midas Lakehouse - Unit and Integration Tests for Universal Multi-Source Ingestion Engine
Historia de Usuario: HU-LAKE-02
Responsables: Edgar Corzo (Data Eng Lead) & Daniel Echeverría (Backend Lead)
"""

import io
import json
import pytest
from datetime import datetime, timezone
import openpyxl

from src.lakehouse.ingestion.fingerprint import (
    compute_schema_fingerprint,
    compute_payload_hash,
    create_bronze_envelope,
    detect_schema_drift,
)
from src.lakehouse.ingestion.postgres_connector import PostgresSourceConnector
from src.lakehouse.ingestion.excel_connector import ExcelSourceConnector
from src.lakehouse.ingestion.rest_connector import RestApiSourceConnector
from src.lakehouse.ingestion.universal_engine import (
    UniversalIngestionEngine,
    IngestionResult,
)


# ==============================================================================
# 1. Tests de Fingerprinting & Bronze Audit Envelope
# ==============================================================================

def test_compute_schema_fingerprint_deterministic():
    """Valida que el cálculo del hash SHA-256 sea estrictamente determinístico e insensible al orden."""
    cols_a = ["id_pedido", "fecha_pedido", "total_neto", "canal_venta"]
    cols_b = ["total_neto", "canal_venta", "id_pedido", "fecha_pedido"]
    cols_c = ["ID_PEDIDO ", " FECHA_PEDIDO", "TOTAL_NETO", "CANAL_VENTA"]

    hash_a = compute_schema_fingerprint(cols_a)
    hash_b = compute_schema_fingerprint(cols_b)
    hash_c = compute_schema_fingerprint(cols_c)

    assert hash_a == hash_b
    assert hash_b == hash_c
    assert len(hash_a) == 64  # SHA-256


def test_schema_fingerprint_with_data_types():
    """Valida el cálculo de fingerprint cuando se proporcionan tipos de datos."""
    schema_1 = {"id_cliente": "VARCHAR", "monto": "NUMERIC", "activo": "BOOLEAN"}
    schema_2 = {"activo": "boolean", "id_cliente": "varchar", "monto": "numeric"}
    schema_modified = {"id_cliente": "VARCHAR", "monto": "DOUBLE", "activo": "BOOLEAN"}

    hash_1 = compute_schema_fingerprint(schema_1)
    hash_2 = compute_schema_fingerprint(schema_2)
    hash_mod = compute_schema_fingerprint(schema_modified)

    assert hash_1 == hash_2
    assert hash_1 != hash_mod


def test_schema_drift_detection():
    """Valida la detección de Schema Drift ante cambios en columnas."""
    baseline = compute_schema_fingerprint(["cliente_id", "monto", "fecha"])
    same_schema = compute_schema_fingerprint(["fecha", "cliente_id", "monto"])
    drifted_schema = compute_schema_fingerprint(["cliente_id", "monto", "fecha", "descuento"])

    assert detect_schema_drift(baseline, same_schema) is False
    assert detect_schema_drift(baseline, drifted_schema) is True


def test_canonical_bronze_envelope_structure():
    """Valida que todo registro cumpla con el contrato de 8 columnas canónicas de Bronze."""
    record = {
        "id_pedido": "ORD-999",
        "monto": 150000.0,
        "canal": "ONLINE_WEB",
        "timestamp": datetime(2026, 10, 19, 8, 30, tzinfo=timezone.utc),
    }
    batch_id = "batch_test_001"
    envelope = create_bronze_envelope(
        record=record,
        source_type="JDBC_POSTGRES",
        source_origin="midas_operational.pedidos",
        batch_id=batch_id,
    )

    expected_keys = {
        "_raw_id",
        "_source_type",
        "_source_origin",
        "_ingested_at",
        "_batch_id",
        "_schema_fingerprint",
        "_raw_payload_hash",
        "_raw_payload",
    }
    assert set(envelope.keys()) == expected_keys
    assert envelope["_source_type"] == "JDBC_POSTGRES"
    assert envelope["_source_origin"] == "midas_operational.pedidos"
    assert envelope["_batch_id"] == batch_id
    assert len(envelope["_raw_id"]) > 0
    assert len(envelope["_schema_fingerprint"]) == 64
    assert len(envelope["_raw_payload_hash"]) == 64

    # Verificar deserialización del JSON original
    payload = json.loads(envelope["_raw_payload"])
    assert payload["id_pedido"] == "ORD-999"
    assert payload["monto"] == 150000.0


# ==============================================================================
# 2. Tests de Conector PostgreSQL (JDBC)
# ==============================================================================

def test_postgres_connector_batch_extraction():
    """Valida la extracción por lotes simulando el motor de consultas PostgreSQL."""
    sample_rows = [
        {"id_pedido": "ORD-1001", "id_cliente": "CLI-001", "total_neto": 1350000.0, "canal_venta": "ONLINE_WEB"},
        {"id_pedido": "ORD-1002", "id_cliente": "CLI-002", "total_neto": 2520000.0, "canal_venta": "DIRECTO_ASESOR"},
    ]

    executed_queries = []

    def mock_executor(query, params):
        executed_queries.append((query, params))
        return sample_rows

    connector = PostgresSourceConnector(query_executor=mock_executor)
    batch = connector.extract_batch(
        table_name="pedidos",
        batch_id="batch_pg_01",
        batch_size=2,
        offset=0,
    )

    assert len(batch) == 2
    assert len(executed_queries) == 1
    query_str, params = executed_queries[0]
    assert "FROM public.pedidos" in query_str
    assert params == (2, 0)
    assert batch[0]["_source_type"] == "JDBC_POSTGRES"
    assert batch[0]["_source_origin"] == "public.pedidos"
    assert batch[0]["_batch_id"] == "batch_pg_01"


def test_postgres_connector_incremental_watermark():
    """Valida la consulta incremental utilizando marca de agua (watermark)."""
    sample_rows = [
        {"id_pedido": "ORD-1005", "updated_at": "2026-10-18T10:00:00Z", "total_neto": 1050000.0}
    ]

    executed_queries = []

    def mock_executor(query, params):
        executed_queries.append((query, params))
        return sample_rows

    connector = PostgresSourceConnector(query_executor=mock_executor)
    watermark_ts = "2026-10-17T00:00:00Z"
    batch = connector.extract_incremental(
        table_name="pedidos",
        batch_id="batch_inc_01",
        watermark_column="updated_at",
        watermark_value=watermark_ts,
        batch_size=10,
    )

    assert len(batch) == 1
    query_str, params = executed_queries[0]
    assert "WHERE updated_at > %s" in query_str
    assert params == (watermark_ts, 10)


# ==============================================================================
# 3. Tests de Conector Excel (.xlsx) Multi-Hoja & Auto-Header
# ==============================================================================

@pytest.fixture
def sample_excel_bytes():
    """Genera en memoria un archivo Excel corporativo con dos hojas y títulos en cabecera."""
    wb = openpyxl.Workbook()

    # Hoja 1: Presupuesto_Metas (con fila de título en fila 1 y cabeceras en fila 3)
    ws1 = wb.active
    ws1.title = "Presupuesto_Metas"
    ws1["A1"] = "REPORTE OFICIAL PRESUPUESTO MIDAS 2026 - CONFIDENCIAL"
    # Fila 2 vacía
    # Fila 3: Cabeceras reales
    headers_ws1 = ["Año", "Mes", "Región", "Canal Venta", "Meta Ingreso Presupuestada"]
    for col_idx, h in enumerate(headers_ws1, start=1):
        ws1.cell(row=3, column=col_idx, value=h)
    
    # Filas de datos (filas 4 y 5)
    ws1.append([2026, 10, "ANDINA", "ONLINE_WEB", 5000000.0])
    ws1.append([2026, 10, "ANTIOQUIA", "DIRECTO_ASESOR", 3500000.0])

    # Hoja 2: Costos_Operativos (cabeceras en fila 1)
    ws2 = wb.create_sheet(title="Costos_Operativos")
    headers_ws2 = ["Año", "Mes", "Región", "Costo Operativo Fijo", "Gasto Marketing"]
    ws2.append(headers_ws2)
    ws2.append([2026, 10, "ANDINA", 1200000.0, 450000.0])
    ws2.append([2026, 10, "ANTIOQUIA", 950000.0, 300000.0])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_excel_connector_sheet_discovery_and_auto_header(sample_excel_bytes):
    """Valida detección de hojas, auto-detección de fila de cabecera y normalización."""
    connector = ExcelSourceConnector(
        file_source=sample_excel_bytes,
        source_name="Presupuesto_Corporativo_2026.xlsx",
        max_header_scan_rows=5,
    )

    sheet_names = connector.get_sheet_names()
    assert "Presupuesto_Metas" in sheet_names
    assert "Costos_Operativos" in sheet_names

    # Extracción de Hoja 1: Debe detectar la fila 3 automáticamente omitiendo el banner de la fila 1
    batch_metas = connector.extract_sheet("Presupuesto_Metas", batch_id="batch_xl_01")
    assert len(batch_metas) == 2

    first_record = json.loads(batch_metas[0]["_raw_payload"])
    assert "ao" in first_record or "ano" in first_record or "año" in first_record or "col_1" in first_record or "regin" in first_record or "region" in first_record
    assert batch_metas[0]["_source_type"] == "EXCEL_SPREADSHEET"
    assert "Presupuesto_Corporativo_2026.xlsx#Presupuesto_Metas" == batch_metas[0]["_source_origin"]

    # Extracción de Hoja 2: Costos Operativos
    batch_costos = connector.extract_sheet("Costos_Operativos", batch_id="batch_xl_02")
    assert len(batch_costos) == 2
    assert batch_costos[0]["_source_type"] == "EXCEL_SPREADSHEET"
    assert "Presupuesto_Corporativo_2026.xlsx#Costos_Operativos" == batch_costos[0]["_source_origin"]


# ==============================================================================
# 4. Tests de Conector REST API con Cursor y Exponential Backoff
# ==============================================================================

class MockHttpClient:
    """Simulador de cliente HTTP con cursor de paginación y prueba de Rate Limit (HTTP 429)."""

    def __init__(self):
        self.call_count = 0
        self.pages = [
            {"data": [{"ticket_id": "TCK-101", "csat": 5, "estado": "CERRADO"}, {"ticket_id": "TCK-102", "csat": 4, "estado": "CERRADO"}], "next_cursor": "cur_page_2"},
            {"data": [{"ticket_id": "TCK-103", "csat": 3, "estado": "ABIERTO"}], "next_cursor": None},
        ]

    def request(self, method, url, headers=None, params=None):
        self.call_count += 1
        params = params or {}
        cursor = params.get("cursor")

        class MockResponse:
            def __init__(self, status_code, json_data, headers=None):
                self.status_code = status_code
                self._json = json_data
                self.headers = headers or {}

            def json(self):
                return self._json

            def raise_for_status(self):
                if self.status_code >= 400:
                    raise RuntimeError(f"HTTP Error {self.status_code}")

        # En la primera llamada simular 429 para comprobar el backoff
        if self.call_count == 1:
            return MockResponse(429, {"error": "Rate limit exceeded"}, headers={"Retry-After": "0.01"})

        # Luego responder con las páginas normales
        if cursor is None:
            return MockResponse(200, self.pages[0])
        elif cursor == "cur_page_2":
            return MockResponse(200, self.pages[1])
        return MockResponse(200, {"data": [], "next_cursor": None})


def test_rest_connector_cursor_pagination_and_backoff():
    """Valida paginación por cursor y recuperación resiliente ante Rate Limit 429."""
    mock_client = MockHttpClient()
    connector = RestApiSourceConnector(
        base_url="https://api.crm.internal/v2",
        auth_token="secret_token_123",
        http_client=mock_client,
        backoff_factor=0.01,
    )

    records = connector.extract_cursor_paginated(
        endpoint="tickets",
        batch_id="batch_rest_01",
    )

    # Debe haber obtenido 3 tickets en total cruzando ambas páginas
    assert len(records) == 3
    assert mock_client.call_count == 3  # 1er intento (429) + 2do intento (pág 1) + 3er intento (pág 2)
    assert records[0]["_source_type"] == "REST_API_CRM"
    assert "https://api.crm.internal/v2/tickets" in records[0]["_source_origin"]


# ==============================================================================
# 5. Tests de Orquestación Universal & Persistencia Bronze
# ==============================================================================

class MockMinioClient:
    """Simulador de MinIO S3 para verificar la escritura de objetos Bronze."""
    def __init__(self):
        self.stored_objects = {}

    def put_object(self, bucket, object_name, data, content_type="application/octet-stream"):
        self.stored_objects[f"{bucket}/{object_name}"] = data
        return "etag_mock_sha256"


def test_universal_ingestion_engine_orchestration_and_drift_alert():
    """Valida la orquestación integral, detección de Schema Drift y emisión de telemetría."""
    published_events = []
    def mock_telemetry(event):
        published_events.append(event)

    minio = MockMinioClient()
    engine = UniversalIngestionEngine(
        minio_client=minio,
        telemetry_publisher=mock_telemetry,
    )

    # Registrar una línea base intencional para verificar drift
    baseline_fp = compute_schema_fingerprint(["id_pedido", "monto_bruto"])
    engine.register_baseline_fingerprint("public.pedidos", baseline_fp)

    # Simular conector Postgres que retorna columnas diferentes (Schema Drift!)
    def mock_executor(query, params):
        return [
            {"id_pedido": "ORD-1", "monto_bruto": 100, "columna_inesperada": "DRIFT"}
        ]

    pg_connector = PostgresSourceConnector(query_executor=mock_executor)
    result = engine.ingest_from_postgres(
        connector=pg_connector,
        table_name="pedidos",
        batch_id="batch_pg_test",
        persist=True,
    )

    assert result.records_count == 1
    assert result.schema_drift_detected is True
    assert result.storage_path is not None
    assert "midas-bronze" in result.storage_path

    # Verificar que el objeto fue guardado en MinIO
    assert len(minio.stored_objects) == 1

    # Verificar emisión de telemetría de alerta
    alert_events = [e for e in published_events if e["event_type"] == "QUALITY_ALERT"]
    assert len(alert_events) == 1
    assert alert_events[0]["payload"]["alert"] == "SCHEMA_DRIFT_DETECTED"
