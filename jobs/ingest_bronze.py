"""Extract a governed source and append its audit envelope to Bronze Parquet using pure PySpark."""

import argparse
import hashlib
import json
import os
import re
import urllib.request
import uuid
from datetime import datetime, timezone

from pyspark.sql import SparkSession, functions as F


def publish(spec: dict, status: str, count: int = 0, detail: str = "") -> None:
    event = {
        "pipeline_id": f"dag_ingest_{spec['conn_id']}",
        "node_id": f"bronze_{spec['conn_id']}",
        "status": status,
        "records_in": count,
        "records_out": count,
        "details": {"event_type": "INGESTION", "message": detail},
    }
    request = urllib.request.Request(
        os.getenv("MIDAS_TELEMETRY_PUBLISH_URL", "http://midas-api:8000/api/v1/telemetry/publish"),
        data=json.dumps(event).encode(),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=10):
            pass
    except Exception:
        # Observabilidad resiliente: no bloquear el job de Spark si la telemetria no responde
        pass


def extract(spark: SparkSession, source: dict):
    kind = source["type"].lower()
    if kind == "postgres":
        url = source.get("connection_url") or os.getenv(source.get("url_env", "POSTGRES_OPERATIONAL_URL"))
        if not url:
            raise ValueError("PostgreSQL connection URL is not set")
        props = {"url": url, "dbtable": f"{source.get('schema', 'public')}.{source['table']}"}
        if source.get("driver"):
            props["driver"] = source["driver"]
        return spark.read.format("jdbc").options(**props).load()

    if kind == "csv":
        return spark.read.option("header", True).option("inferSchema", True).csv(source["path"])

    if kind == "excel":
        from openpyxl import load_workbook

        wb = load_workbook(source["path"], read_only=True, data_only=True)
        sheet_name = source.get("sheet_name") or wb.sheetnames[0]
        sheet = wb[sheet_name]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            wb.close()
            return spark.createDataFrame([], schema="")

        headers = [
            re.sub(r"[^\w\s]", "", str(c).strip().lower()).replace(" ", "_") if c is not None else f"col_{i}"
            for i, c in enumerate(header_row)
        ]
        data_rows = []
        for r in rows_iter:
            if any(cell is not None for cell in r):
                row_dict = {
                    headers[i]: (str(cell) if cell is not None else None)
                    for i, cell in enumerate(r)
                    if i < len(headers)
                }
                data_rows.append(row_dict)
        wb.close()
        return spark.createDataFrame(data_rows)

    if kind == "rest_api":
        target_url = source.get("base_url") or source.get("url") or source.get("endpoint_url")
        if not target_url:
            raise ValueError("REST API target URL is not set")
        headers = json.loads(os.environ[source["headers_env"]]) if source.get("headers_env") else {}
        request = urllib.request.Request(target_url, headers=headers)
        with urllib.request.urlopen(request, timeout=int(source.get("timeout_seconds", 30))) as response:
            payload = json.loads(response.read())
        rows = payload.get(source.get("data_key", "data"), payload) if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            raise ValueError("REST response must contain a list of records")
        return spark.createDataFrame(rows)

    raise ValueError(f"Unsupported source type: {kind}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec-json", required=True)
    spec = json.loads(parser.parse_args().spec_json)
    source = spec["source"]
    source_name = source.get("table") or source.get("path") or source.get("url") or spec["conn_id"]
    dataset = re.sub(r"[^A-Za-z0-9_-]", "_", str(source_name).rstrip("/").split("/")[-1])
    batch_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc)
    output = spec.get("bronze", {}).get("path") or (
        f"s3a://midas-bronze/sources/{spec['conn_id']}/{dataset}/year={now:%Y}/"
        f"month={now:%m}/day={now:%d}/{batch_id}.parquet"
    )

    spark = (
        SparkSession.builder.appName(f"midas-ingest-{spec['conn_id']}")
        .config("spark.sql.extensions", "org.apache.iceberg.spark.extensions.IcebergSparkSessionExtensions")
        .config("spark.sql.catalog.iceberg", "org.apache.iceberg.spark.SparkCatalog")
        .config("spark.sql.catalog.iceberg.type", "hadoop")
        .config("spark.sql.catalog.iceberg.warehouse", os.getenv("ICEBERG_WAREHOUSE", "s3a://midas-lakehouse/iceberg"))
        .config("spark.hadoop.fs.s3a.endpoint", os.getenv("MINIO_ENDPOINT", "http://minio:9000"))
        .config("spark.hadoop.fs.s3a.access.key", os.getenv("MINIO_ROOT_USER", "midas_admin"))
        .config("spark.hadoop.fs.s3a.secret.key", os.getenv("MINIO_ROOT_PASSWORD", "midas_storage_2026"))
        .config("spark.hadoop.fs.s3a.path.style.access", "true")
        .config("spark.hadoop.fs.s3a.impl", "org.apache.hadoop.fs.s3a.S3AFileSystem")
        .config("spark.sql.shuffle.partitions", "4")
        .getOrCreate()
    )

    count = 0
    publish(spec, "RUNNING", detail="ingestion started")
    try:
        rows = extract(spark, source)
        columns = rows.columns
        count = rows.count()
        payload = F.to_json(F.struct(*[F.col(name) for name in columns]))

        source_type_map = {
            "postgres": "JDBC_POSTGRES",
            "csv": "CSV_STORAGE",
            "excel": "EXCEL_SPREADSHEET",
            "rest_api": "REST_API",
        }
        stype_label = source_type_map.get(source["type"].lower(), "GENERIC_INGESTION")

        enriched = (
            rows.withColumn("_raw_id", F.expr("uuid()"))
            .withColumn("_source_type", F.lit(stype_label))
            .withColumn("_source_origin", F.lit(source.get("table", source.get("path", source.get("url", spec["conn_id"])))))
            .withColumn("_ingested_at", F.current_timestamp())
            .withColumn("_batch_id", F.lit(batch_id))
            .withColumn("_schema_fingerprint", F.lit(hashlib.sha256("|".join(sorted(columns)).encode()).hexdigest()))
            .withColumn("_raw_payload", payload)
            .withColumn("_raw_payload_hash", F.sha2(payload, 256))
        )
        enriched.write.mode("append").parquet(output)
        publish(spec, "HEALTHY", count, f"Bronze parquet written to {output}")
    except Exception as exc:
        publish(spec, "ERROR", count, str(exc)[:500])
        raise
    finally:
        spark.stop()


if __name__ == "__main__":
    main()
